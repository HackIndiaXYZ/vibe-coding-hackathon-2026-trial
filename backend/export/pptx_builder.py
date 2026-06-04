from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


# ── Palette ──────────────────────────────────────────────────────────────────
C_BG        = RGBColor(0x0F, 0x17, 0x2A)   # deep navy
C_ACCENT    = RGBColor(0x38, 0xBD, 0xF8)   # sky blue
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT     = RGBColor(0xCB, 0xD5, 0xE1)   # slate-300
C_BULLET_BG = RGBColor(0x1E, 0x2D, 0x45)   # slightly lighter navy
C_NOTE_BG   = RGBColor(0x1A, 0x28, 0x40)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

BULLET_ICONS = ["▶", "◆", "●", "★", "→"]


def _rgb(shape, color: RGBColor):
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_textbox(slide, left, top, width, height,
                 text, size, bold=False, color=C_WHITE,
                 align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(left, top, width, height)
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return txb


def _set_bg(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


# ── Title slide ───────────────────────────────────────────────────────────────
def _make_title_slide(prs, topic, grade, subject):
    layout = prs.slide_layouts[6]          # blank
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, C_BG)

    # accent bar left edge
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.18), SLIDE_H)
    _rgb(bar, C_ACCENT)

    # top label
    _add_textbox(slide,
                 Inches(0.5), Inches(1.8), Inches(12), Inches(0.6),
                 f"{subject}  ·  {grade}",
                 Pt(18), color=C_ACCENT)

    # main title
    _add_textbox(slide,
                 Inches(0.5), Inches(2.4), Inches(12), Inches(2),
                 topic,
                 Pt(52), bold=True, color=C_WHITE)

    # divider line
    line = slide.shapes.add_shape(1,
                                  Inches(0.5), Inches(4.55),
                                  Inches(5), Inches(0.05))
    _rgb(line, C_ACCENT)

    # subtitle
    _add_textbox(slide,
                 Inches(0.5), Inches(4.7), Inches(10), Inches(0.5),
                 "AI-Generated Lesson Pack",
                 Pt(16), color=C_LIGHT)


# ── Content slide ─────────────────────────────────────────────────────────────
def _make_content_slide(prs, slide_data: dict):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    _set_bg(slide, C_BG)

    number      = slide_data.get("slide_number", "")
    title       = slide_data.get("title", "")
    key_points  = slide_data.get("key_points", [])
    body        = slide_data.get("body", "")
    examples    = slide_data.get("examples", [])
    discussion  = slide_data.get("discussion_question", "")
    visual      = slide_data.get("visual_suggestion", "")
    note        = slide_data.get("teacher_note", "")

    # accent bar
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.18), SLIDE_H)
    _rgb(bar, C_ACCENT)

    # slide number chip
    chip = slide.shapes.add_shape(1,
                                  Inches(0.35), Inches(0.28),
                                  Inches(0.55), Inches(0.38))
    _rgb(chip, C_ACCENT)
    _add_textbox(slide,
                 Inches(0.36), Inches(0.28), Inches(0.55), Inches(0.38),
                 str(number), Pt(13), bold=True, color=C_BG,
                 align=PP_ALIGN.CENTER)

    # title
    _add_textbox(slide,
                 Inches(1.1), Inches(0.22), Inches(11.7), Inches(0.7),
                 title, Pt(28), bold=True, color=C_WHITE)

    # accent underline
    ul = slide.shapes.add_shape(1,
                                Inches(1.1), Inches(0.95),
                                Inches(11.3), Inches(0.04))
    _rgb(ul, C_ACCENT)

    # ── bullet cards ─────────────────────────────────────────────────────────
    card_top    = Inches(1.15)
    card_left   = Inches(0.45)
    card_w      = Inches(7.9)
    card_h      = Inches(0.72)
    card_gap    = Inches(0.12)

    for i, point in enumerate(key_points[:6]):
        top = card_top + i * (card_h + card_gap)

        # card background
        card = slide.shapes.add_shape(1, card_left, top, card_w, card_h)
        _rgb(card, C_BULLET_BG)
        card.line.color.rgb = C_ACCENT
        card.line.width = Emu(12700)          # 1 pt

        # icon
        icon = BULLET_ICONS[i % len(BULLET_ICONS)]
        _add_textbox(slide,
                     card_left + Inches(0.1), top + Inches(0.08),
                     Inches(0.45), card_h,
                     icon, Pt(14), color=C_ACCENT)

        # point text
        _add_textbox(slide,
                     card_left + Inches(0.55), top + Inches(0.07),
                     card_w - Inches(0.65), card_h - Inches(0.1),
                     point, Pt(15), color=C_WHITE)

    # ── body + examples below bullets ────────────────────────────────────────
    body_top = card_top + len(key_points[:6]) * (card_h + card_gap) + Inches(0.1)

    if body:
        _add_textbox(slide,
                     card_left, body_top, card_w, Inches(0.9),
                     body, Pt(12), color=C_LIGHT)
        body_top += Inches(1.0)

    if examples:
        _add_textbox(slide,
                     card_left, body_top, card_w, Inches(0.25),
                     "EXAMPLES", Pt(10), bold=True, color=C_ACCENT)
        body_top += Inches(0.28)
        for ex in examples[:3]:
            _add_textbox(slide,
                         card_left + Inches(0.15), body_top, card_w - Inches(0.15), Inches(0.28),
                         f"• {ex}", Pt(11), color=C_LIGHT)
            body_top += Inches(0.3)

    if discussion:
        body_top += Inches(0.08)
        dbox = slide.shapes.add_shape(1, card_left, body_top, card_w, Inches(0.5))
        _rgb(dbox, RGBColor(0x1E, 0x3A, 0x5F))
        dbox.line.color.rgb = RGBColor(0x60, 0xA5, 0xFA)
        dbox.line.width = Emu(12700)
        _add_textbox(slide,
                     card_left + Inches(0.12), body_top + Inches(0.05),
                     card_w - Inches(0.2), Inches(0.45),
                     f"💬 {discussion}", Pt(11), color=RGBColor(0xBF, 0xDB, 0xFE))

    # ── teacher note panel (right column) ────────────────────────────────────
    if note:
        panel_left = Inches(8.6)
        panel_top  = Inches(1.15)
        panel_w    = Inches(4.5)
        panel_h    = Inches(5.6)

        panel = slide.shapes.add_shape(1, panel_left, panel_top, panel_w, panel_h)
        _rgb(panel, C_NOTE_BG)
        panel.line.color.rgb = C_ACCENT
        panel.line.width = Emu(12700)

        _add_textbox(slide,
                     panel_left + Inches(0.15), panel_top + Inches(0.12),
                     panel_w - Inches(0.3), Inches(0.35),
                     "TEACHER NOTE", Pt(11), bold=True, color=C_ACCENT)

        _add_textbox(slide,
                     panel_left + Inches(0.15), panel_top + Inches(0.5),
                     panel_w - Inches(0.3), panel_h - Inches(0.65),
                     note, Pt(13), color=C_LIGHT)

    # ── slide notes (PowerPoint notes pane) ──────────────────────────────────
    notes_parts = []
    if note:        notes_parts.append(f"Teacher note: {note}")
    if visual:      notes_parts.append(f"Visual: {visual}")
    if discussion:  notes_parts.append(f"Discussion: {discussion}")
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "\n\n".join(notes_parts)


# ── Public entry point ────────────────────────────────────────────────────────
def build_pptx(lesson: dict) -> BytesIO:
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    _make_title_slide(prs, lesson["topic"], lesson["grade"], lesson["subject"])

    for slide_data in lesson.get("slides", []):
        _make_content_slide(prs, slide_data)

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
